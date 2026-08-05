"""
Build the comparison workbook of the parameter study: for every run of
params_study.csv, overlay the pyFBS (nonlinearFBS, ModalVPFRF) and pyhbm
(RBE2+CB+HBM reference) forced-response branches, compute purely DESCRIPTIVE
agreement metrics (no root-cause analysis) and assemble everything into
comparison.xlsx:

  sheet "Uebersicht"   parameter matrix + metrics per run, traffic-light fills
  one sheet per run    embedded overlay plot (log scale) + parameter/metric block

Metrics per run (relative to the pyhbm reference):
  peak amplitude / peak frequency of each branch and their relative deviation;
  envelope deviation: common 10-Hz bins, envelope = max(amp) per bin (robust
  against the multivalued fold regions of an NFRC), median and maximum of
  |env_pyfbs - env_pyhbm| / env_pyhbm over bins covered by BOTH solvers.

Traffic-light thresholds (TRESH_*): green < 5 %, yellow < 20 %, red otherwise.
"""

import csv
import datetime
from pathlib import Path

import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import Workbook, load_workbook
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.chart.marker import Marker
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
RESULTS = HERE / "results"
PLOTS = RESULTS / "plots"
PARAMS_CSV = HERE / "params_study.csv"
XLSX_OUT = HERE / "comparison.xlsx"

F0 = 50.0
F_LO, F_HI = 20.0, 1000.0
BIN_WIDTH = 10.0                       # [Hz] envelope bins
THRESH_GREEN, THRESH_YELLOW = 5.0, 20.0   # [%] traffic-light thresholds
DATA_COL = 18                          # raw curve columns start here (right of PNG)

FILL_GREEN = PatternFill("solid", start_color="C6EFCE")
FILL_YELLOW = PatternFill("solid", start_color="FFEB9C")
FILL_RED = PatternFill("solid", start_color="FFC7CE")
BOLD = Font(bold=True)


# ---------------------------------------------------------------------------
# loading
# ---------------------------------------------------------------------------

def read_params():
    with open(PARAMS_CSV, newline="", encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def load_branch(run_id, solver):
    """(freq, amp) of one solver branch, or None if the run has no result."""
    path = RESULTS / f"{run_id}_{solver}.csv"
    if not path.exists():
        return None
    d = np.loadtxt(path, delimiter=",", skiprows=1, ndmin=2)
    if d.size == 0:
        return None
    return d[:, 0], d[:, 2]


def load_linear(run_id, solver):
    path = RESULTS / f"{run_id}_linear_{solver}.csv"
    if not path.exists():
        return None
    d = np.loadtxt(path, delimiter=",", skiprows=1, ndmin=2)
    return d[:, 0], d[:, 1]


def load_manifests():
    """Last manifest line per (run_id, solver): status, n_points, runtime_s."""
    info = {}
    for solver in ("pyfbs", "pyhbm"):
        path = RESULTS / f"manifest_{solver}.csv"
        if not path.exists():
            continue
        with open(path, newline="", encoding="utf-8") as fh:
            for row in csv.DictReader(fh):
                info[row["run_id"], solver] = row
    return info


# ---------------------------------------------------------------------------
# metrics
# ---------------------------------------------------------------------------

def envelope(freq, amp, edges):
    """Per-bin maximum amplitude (NaN where the branch has no points)."""
    env = np.full(len(edges) - 1, np.nan)
    idx = np.digitize(freq, edges) - 1
    for b in np.unique(idx):
        if 0 <= b < len(env):
            env[b] = amp[idx == b].max()
    return env


def run_metrics(fbs, hbm):
    """Descriptive agreement numbers of one run; None entries where a branch
    is missing. Deviations are relative to the pyhbm reference, in percent."""
    m = {}
    edges = np.arange(F_LO, F_HI + BIN_WIDTH, BIN_WIDTH)
    centers = 0.5 * (edges[:-1] + edges[1:])
    for name, cur in (("pyfbs", fbs), ("pyhbm", hbm)):
        if cur is not None:
            i = int(np.argmax(cur[1]))
            m[f"peak_amp_{name}"] = float(cur[1][i])
            m[f"peak_freq_{name}"] = float(cur[0][i])
    if fbs is not None and hbm is not None:
        m["dev_peak_amp"] = 100.0 * abs(m["peak_amp_pyfbs"] - m["peak_amp_pyhbm"]) / m["peak_amp_pyhbm"]
        m["dev_peak_freq"] = 100.0 * abs(m["peak_freq_pyfbs"] - m["peak_freq_pyhbm"]) / m["peak_freq_pyhbm"]
        env_f = envelope(*fbs, edges)
        env_h = envelope(*hbm, edges)
        both = ~np.isnan(env_f) & ~np.isnan(env_h)
        dev = np.full(len(centers), np.nan)
        dev[both] = 100.0 * np.abs(env_f[both] - env_h[both]) / env_h[both]
        m["dev_env_median"] = float(np.median(dev[both])) if both.any() else np.nan
        m["dev_env_max"] = float(np.max(dev[both])) if both.any() else np.nan
        m["bins_single"] = int(np.count_nonzero(np.isnan(env_f) != np.isnan(env_h)))
        m["dev_curve"] = (centers, dev)
    return m


# ---------------------------------------------------------------------------
# plots
# ---------------------------------------------------------------------------

def render_plot(run_id, comment, fbs, hbm, lin_f, lin_h, metrics):
    """Two panels: NFRC overlay (log y) and relative envelope deviation."""
    fig, (ax, axd) = plt.subplots(
        2, 1, figsize=(9.0, 7.0), sharex=True,
        gridspec_kw={"height_ratios": [3, 1]})

    if lin_f is not None:
        ax.semilogy(lin_f[0], lin_f[1] * F0, "-", color="#999999", lw=1.0,
                    label=f"linear (pyFBS LM-FBS) x F0={F0:g} N")
    if lin_h is not None:
        ax.semilogy(lin_h[0], lin_h[1] * F0, "--", color="#555555", lw=1.0,
                    label=f"linear (pyhbm RBE2+CB) x F0={F0:g} N")
    if hbm is not None:
        ax.semilogy(hbm[0], hbm[1], "-", color="#d62728", lw=1.4,
                    label="pyhbm RBE2+CB+HBM (Referenz)")
    if fbs is not None:
        ax.semilogy(fbs[0], fbs[1], "-", color="#2ca02c", lw=1.4,
                    label="pyFBS nonlinearFBS (ModalVPFRF)")
    ax.set_xlim(F_LO, F_HI)
    ax.set_ylabel("Amplitude |u_out|  [m]")
    ax.set_title(f"{run_id}: {comment}")
    ax.grid(True, which="both", alpha=0.3)
    ax.legend(fontsize=8)

    if "dev_curve" in metrics:
        centers, dev = metrics["dev_curve"]
        axd.plot(centers, dev, "-", color="#1f77b4", lw=1.0)
        axd.axhline(THRESH_GREEN, color="#2ca02c", lw=0.8, ls=":")
        axd.axhline(THRESH_YELLOW, color="#d62728", lw=0.8, ls=":")
        axd.set_ylim(bottom=0)
    else:
        axd.text(0.5, 0.5, "kein Vergleich moeglich (Lauf fehlt)",
                 ha="center", va="center", transform=axd.transAxes)
    axd.set_xlabel("Frequenz [Hz]")
    axd.set_ylabel("Huellkurven-\nAbw. [%]")
    axd.grid(True, alpha=0.3)

    fig.tight_layout()
    PLOTS.mkdir(exist_ok=True)
    path = PLOTS / f"{run_id}.png"
    fig.savefig(path, dpi=110)
    plt.close(fig)
    return path


# ---------------------------------------------------------------------------
# workbook
# ---------------------------------------------------------------------------

PARAM_COLS = ["k_trans", "k_rot", "alpha_trans", "alpha_rot",
              "beta_trans", "beta_rot"]
METRIC_COLS = [
    ("peak_amp_pyfbs", "Peak-Amp pyFBS [m]", "0.00E+00", False),
    ("peak_amp_pyhbm", "Peak-Amp pyhbm [m]", "0.00E+00", False),
    ("dev_peak_amp", "Abw. Peak-Amp [%]", "0.0", True),
    ("peak_freq_pyfbs", "Peak-Freq pyFBS [Hz]", "0.0", False),
    ("peak_freq_pyhbm", "Peak-Freq pyhbm [Hz]", "0.0", False),
    ("dev_peak_freq", "Abw. Peak-Freq [%]", "0.0", True),
    ("dev_env_median", "Huellkurven-Abw. Median [%]", "0.0", True),
    ("dev_env_max", "Huellkurven-Abw. Max [%]", "0.0", True),
    ("bins_single", "Bins nur 1 Solver", "0", False),
]
LEGEND = [
    "Legende:",
    f"amp_m = Spitzenwert |u_out(t)| ueber eine Periode am Output-DoF (S1 X auf A); Anregung F0={F0:g} N (H28 auf B); Sweep 1000->20 Hz, Harmonische 1,3,5,7.",
    "Kurven liegen in Branch-Reihenfolge der Kontinuation vor (NFRC kann mehrdeutig sein); Details je Lauf auf dem gleichnamigen Blatt.",
    f"Huellkurve = max(amp_m) je {BIN_WIDTH:g}-Hz-Bin (20-1000 Hz); Abweichungen relativ zur pyhbm-Referenz; 'Bins nur 1 Solver' = Bins, die nur einer der beiden Aeste abdeckt.",
    f"Ampel: gruen < {THRESH_GREEN:g} %, gelb {THRESH_GREEN:g}-{THRESH_YELLOW:g} %, rot > {THRESH_YELLOW:g} %.",
    "beta-Werte der D-Laeufe sind aus B0 kalibriert (Daempfkraft im Peak ca. 10 % / 50 % der linearen Interface-Kraftskala, getrennt trans/rot); siehe calibrate_beta.py.",
    "Nur deskriptiver Vergleich beider Solver -- keine Ursachenanalyse.",
]


def fill_for(value):
    if value is None or (isinstance(value, float) and np.isnan(value)):
        return None
    if value < THRESH_GREEN:
        return FILL_GREEN
    if value < THRESH_YELLOW:
        return FILL_YELLOW
    return FILL_RED


def write_curve(ws, col, header_x, header_y, x, y, row0=1):
    """Write an (x, y) curve into columns col / col+1 (headers in row0) and
    return (xref, yref) for a chart Series. The numbers stay inspectable next
    to the chart."""
    ws.cell(row0, col, header_x)
    ws.cell(row0, col + 1, header_y)
    for i, (xi, yi) in enumerate(zip(x, y), start=row0 + 1):
        ws.cell(i, col, float(xi))
        ws.cell(i, col + 1, float(yi))
    last = row0 + len(x)
    return (Reference(ws, min_col=col, min_row=row0 + 1, max_row=last),
            Reference(ws, min_col=col + 1, min_row=row0 + 1, max_row=last))


def line_series(refs, title, color, width_pt=1.5, dashed=False):
    """A line-only (marker-free) scatter Series from a (xref, yref) pair."""
    xref, yref = refs
    s = Series(yref, xref, title=title)               # (values, xvalues)
    s.marker = Marker(symbol="none")
    s.graphicalProperties.line.solidFill = color
    s.graphicalProperties.line.width = int(width_pt * 12700)   # pt -> EMU
    if dashed:
        s.graphicalProperties.line.dashStyle = "dash"
    s.smooth = False
    return s


def add_run_charts(ws, run_id, row, metrics, data):
    """Native (zoomable, vector) scatter charts below the PNG: the NFRC overlay
    on a log amplitude axis plus the relative envelope-deviation curve. Points
    are connected in branch order, exactly like the PNG (the NFRC is
    multivalued, so the fold is traced, not sorted by frequency)."""
    col = DATA_COL
    overlay = ScatterChart()
    overlay.title = f"{run_id}: {row['comment']}"
    overlay.x_axis.title = "Frequenz [Hz]"
    overlay.y_axis.title = "Amplitude |u_out| [m]"
    overlay.x_axis.scaling.min, overlay.x_axis.scaling.max = F_LO, F_HI
    overlay.y_axis.scaling.logBase = 10
    overlay.x_axis.delete = overlay.y_axis.delete = False
    overlay.height, overlay.width = 12, 26            # cm

    for key, hdr, color, dashed, scale, width in (
            ("lin_h", "lin pyhbm x F0", "888888", True, F0, 1.0),
            ("lin_f", "lin pyFBS x F0", "BBBBBB", True, F0, 1.0),
            ("hbm", "pyhbm (Referenz)", "D62728", False, 1.0, 1.6),
            ("fbs", "pyFBS (ModalVPFRF)", "2CA02C", False, 1.0, 1.6)):
        cur = data[key]
        if cur is None:
            continue
        refs = write_curve(ws, col, f"{key}_freq_hz", hdr, cur[0], cur[1] * scale)
        col += 3                                      # blank spacer column between blocks
        overlay.series.append(line_series(refs, hdr, color, width, dashed))
    ws.add_chart(overlay, "A56")

    if "dev_curve" in metrics:
        centers, dev = metrics["dev_curve"]
        ok = np.isfinite(dev)
        if ok.any():
            refs = write_curve(ws, col, "dev_freq_hz", "dev_prozent",
                               centers[ok], dev[ok])
            devc = ScatterChart()
            devc.title = "Huellkurven-Abweichung [%] (relativ zu pyhbm)"
            devc.x_axis.title = "Frequenz [Hz]"
            devc.y_axis.title = "Abweichung [%]"
            devc.x_axis.scaling.min, devc.x_axis.scaling.max = F_LO, F_HI
            devc.x_axis.delete = devc.y_axis.delete = False
            devc.height, devc.width = 7, 26
            devc.series.append(line_series(refs, "Abw. [%]", "1F77B4"))
            ws.add_chart(devc, "A82")


def build_workbook(rows, all_metrics, all_data, manifests):
    wb = Workbook()
    ws = wb.active
    ws.title = "Uebersicht"

    headers = (["Lauf", "Kommentar"] + PARAM_COLS
               + ["Status pyFBS", "Status pyhbm", "Punkte pyFBS", "Punkte pyhbm",
                  "Laufzeit pyFBS [s]", "Laufzeit pyhbm [s]"]
               + [h for _, h, _, _ in METRIC_COLS])
    ws.append(headers)
    for c in ws[1]:
        c.font = BOLD
        c.alignment = Alignment(wrap_text=True, vertical="top")

    for row in rows:
        run_id = row["run_id"]
        m = all_metrics[run_id]
        line = [run_id, row["comment"]]
        line += [float(row[p]) for p in PARAM_COLS]
        for solver in ("pyfbs", "pyhbm"):
            info = manifests.get((run_id, solver))
            line.append(info["status"] if info else "-")
        for solver in ("pyfbs", "pyhbm"):
            info = manifests.get((run_id, solver))
            line.append(int(info["n_points"]) if info else None)
        for solver in ("pyfbs", "pyhbm"):
            info = manifests.get((run_id, solver))
            line.append(float(info["runtime_s"]) if info else None)
        for key, _, _, _ in METRIC_COLS:
            v = m.get(key)
            line.append(None if v is None or (isinstance(v, float) and np.isnan(v)) else v)
        ws.append(line)

        r = ws.max_row
        ws.cell(r, 1).hyperlink = f"#'{run_id}'!A1"
        ws.cell(r, 1).font = Font(color="0563C1", underline="single")
        for j, p in enumerate(PARAM_COLS, start=3):
            ws.cell(r, j).number_format = "0.0E+00"
        base = 2 + len(PARAM_COLS) + 6
        for j, (key, _, fmt, ampel) in enumerate(METRIC_COLS, start=base + 1):
            cell = ws.cell(r, j)
            cell.number_format = fmt
            if ampel:
                f = fill_for(m.get(key))
                if f is not None:
                    cell.fill = f

    ws.freeze_panes = "C2"
    widths = [6, 28] + [10] * len(PARAM_COLS) + [11, 11, 8, 8, 9, 9] + [13] * len(METRIC_COLS)
    for j, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w

    r0 = ws.max_row + 2
    for i, text in enumerate(LEGEND):
        cell = ws.cell(r0 + i, 1, text)
        if i == 0:
            cell.font = BOLD
    ws.cell(r0 + len(LEGEND) + 1, 1,
            f"Erstellt {datetime.date.today().isoformat()} aus results/*.csv "
            f"(build_comparison.py)")

    # one sheet per run: parameter/metric block + embedded plot
    for row in rows:
        run_id = row["run_id"]
        m = all_metrics[run_id]
        s = wb.create_sheet(run_id)
        s["A1"] = f"Lauf {run_id}: {row['comment']}"
        s["A1"].font = BOLD
        s["A2"] = "zurueck zur Uebersicht"
        s["A2"].hyperlink = "#'Uebersicht'!A1"
        s["A2"].font = Font(color="0563C1", underline="single")

        s["A4"] = "Parameter"
        s["A4"].font = BOLD
        for i, p in enumerate(PARAM_COLS):
            s.cell(5 + i, 1, p)
            s.cell(5 + i, 2, float(row[p])).number_format = "0.0E+00"

        s["D4"] = "Kennzahlen (relativ zu pyhbm)"
        s["D4"].font = BOLD
        for i, (key, label, fmt, ampel) in enumerate(METRIC_COLS):
            s.cell(5 + i, 4, label)
            v = m.get(key)
            if v is not None and not (isinstance(v, float) and np.isnan(v)):
                cell = s.cell(5 + i, 6, v)
                cell.number_format = fmt
                if ampel:
                    f = fill_for(v)
                    if f is not None:
                        cell.fill = f
            else:
                s.cell(5 + i, 6, "-")
        s.column_dimensions["A"].width = 14
        s.column_dimensions["D"].width = 26

        png = PLOTS / f"{run_id}.png"
        if png.exists():
            img = XLImage(str(png))
            s.add_image(img, "A16")

        # zoomable native charts below the PNG + raw curve data to the right
        add_run_charts(s, run_id, row, m, all_data[run_id])

    wb.save(XLSX_OUT)
    print(f"workbook written: {XLSX_OUT}")


def main():
    global RESULTS, PLOTS, PARAMS_CSV, XLSX_OUT, F_LO, F_HI
    import argparse
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--params", help="params CSV filename (default: params_study.csv)")
    ap.add_argument("--results", help="results subfolder name (default: results)")
    ap.add_argument("--out", help="output workbook filename (default: comparison.xlsx)")
    ap.add_argument("--flo", type=float, help="plot/bin lower bound [Hz] (default: 20)")
    ap.add_argument("--fhi", type=float, help="plot/bin upper bound [Hz] (default: 1000)")
    args = ap.parse_args()
    if args.params:
        PARAMS_CSV = HERE / args.params
    if args.results:
        RESULTS = HERE / args.results
        PLOTS = RESULTS / "plots"
    if args.out:
        XLSX_OUT = HERE / args.out
    if args.flo is not None:
        F_LO = args.flo
    if args.fhi is not None:
        F_HI = args.fhi

    rows = read_params()
    manifests = load_manifests()
    all_metrics, all_data = {}, {}
    for row in rows:
        run_id = row["run_id"]
        fbs = load_branch(run_id, "pyfbs")
        hbm = load_branch(run_id, "pyhbm")
        lin_f = load_linear(run_id, "pyfbs")
        lin_h = load_linear(run_id, "pyhbm")
        m = run_metrics(fbs, hbm)
        all_metrics[run_id] = m
        all_data[run_id] = dict(fbs=fbs, hbm=hbm, lin_f=lin_f, lin_h=lin_h)
        render_plot(run_id, row["comment"], fbs, hbm, lin_f, lin_h, m)
        state = "ok" if "dev_env_median" in m else "unvollstaendig"
        print(f"[{run_id}] {state}")

    build_workbook(rows, all_metrics, all_data, manifests)

    load_workbook(XLSX_OUT)                     # reload check: file opens cleanly
    print(f"reload check ok ({XLSX_OUT.stat().st_size / 1024:.0f} kB)")


if __name__ == "__main__":
    main()
